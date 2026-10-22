#!/usr/bin/env python3
import argparse
import csv
import json
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

CUTOFF = date(2024, 12, 31)
PAYMENTS_FILE = "Chicago_Payments_2024.csv"
CONTRACTS_FILE = "Chicago_Contracts_2019plus.csv"
DIRECT_MARKERS = {"", "DV"}
MONEY_FMT = '$#,##0.00;[Red]-$#,##0.00'
RATIO_FMT = '0.00x'


def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return reader.fieldnames, [(row_number, row) for row_number, row in enumerate(reader, 2)]


def money(value):
    return Decimal(value) if value not in (None, "") else None


def iso_date(value):
    return datetime.fromisoformat(value).date()


def decimal_text(value):
    if value is None:
        return None
    return format(value, "f")


def payment_rows(path):
    fields, raw = read_csv(path)
    required = {"voucher_number", "amount", "check_date", "department_name", "contract_number", "vendor_name"}
    if set(fields) != required:
        raise ValueError(f"付款字段与预期不符: {fields}")
    rows = []
    for source_row, row in raw:
        rows.append({**row, "source_row": source_row, "amount_decimal": money(row["amount"])})
    return fields, rows


def contract_rows(path):
    fields, raw = read_csv(path)
    required = {
        "purchase_order_description", "purchase_order_contract_number", "revision_number",
        "specification_number", "contract_type", "start_date", "end_date", "approval_date",
        "department", "vendor_name", "vendor_id", "address_1", "address_2", "city", "state",
        "zip", "award_amount", "procurement_type", "contract_pdf"
    }
    if set(fields) != required:
        raise ValueError(f"合同字段与预期不符: {fields}")
    rows = []
    for source_row, row in raw:
        rows.append({
            **row,
            "source_row": source_row,
            "approval_date_value": iso_date(row["approval_date"]),
            "revision_value": int(row["revision_number"]),
            "award_decimal": money(row["award_amount"]),
        })
    return fields, rows


def exact_tuple(row, fields):
    return tuple(row[field] for field in fields)


def duplicate_stats(rows, fields):
    counts = Counter(exact_tuple(row, fields) for row in rows)
    return {
        "duplicate_groups": sum(count > 1 for count in counts.values()),
        "duplicate_extra_rows": sum(count - 1 for count in counts.values() if count > 1),
        "maximum_multiplicity": max(counts.values(), default=0),
    }


def select_contract_view(rows, cutoff=None):
    grouped = defaultdict(list)
    for row in rows:
        if cutoff is None or row["approval_date_value"] <= cutoff:
            grouped[row["purchase_order_contract_number"]].append(row)
    selected = {}
    for contract_number, candidates in grouped.items():
        latest_approval = max(row["approval_date_value"] for row in candidates)
        candidates = [row for row in candidates if row["approval_date_value"] == latest_approval]
        highest_revision = max(row["revision_value"] for row in candidates)
        selected[contract_number] = [row for row in candidates if row["revision_value"] == highest_revision]
    return selected


def payment_status(row, year_end_view):
    contract_number = row["contract_number"]
    if contract_number in DIRECT_MARKERS:
        return "直接付款/空号"
    if contract_number not in year_end_view:
        return "未匹配"
    if len(year_end_view[contract_number]) != 1:
        return "版本冲突"
    return "唯一匹配"


def summarize_status(rows, view):
    buckets = defaultdict(list)
    for row in rows:
        buckets[payment_status(row, view)].append(row)
    order = ["唯一匹配", "未匹配", "版本冲突", "直接付款/空号"]
    return [{
        "status": status,
        "row_count": len(buckets[status]),
        "contract_count": len({row["contract_number"] for row in buckets[status] if row["contract_number"] not in DIRECT_MARKERS}),
        "amount": sum((row["amount_decimal"] for row in buckets[status]), Decimal("0")),
    } for status in order]


def aggregate_annual(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["contract_number"]].append(row)
    return grouped


def source_rows_text(rows):
    return ", ".join(str(row["source_row"]) for row in sorted(rows, key=lambda item: item["source_row"]))


def vendors_text(rows):
    return " | ".join(sorted({row["vendor_name"] for row in rows}))


def build_analysis(payment_fields, payments, contract_fields, contracts):
    daily = [row for row in payments if row["check_date"].count("/") == 2]
    annual = [row for row in payments if row["check_date"] == "2024"]
    unexplained_dates = [row for row in payments if row not in daily and row not in annual]
    year_end_view = select_contract_view(contracts, CUTOFF)
    full_latest_view = select_contract_view(contracts)
    annual_by_contract = aggregate_annual(annual)

    daily_contracts = {row["contract_number"] for row in daily if row["contract_number"] not in DIRECT_MARKERS}
    annual_contracts = {row["contract_number"] for row in annual if row["contract_number"] not in DIRECT_MARKERS}
    overlap = []
    for contract_number in sorted(daily_contracts & annual_contracts):
        daily_rows = [row for row in daily if row["contract_number"] == contract_number]
        annual_rows = annual_by_contract[contract_number]
        overlap.append({
            "contract_number": contract_number,
            "daily_source_rows": source_rows_text(daily_rows),
            "daily_row_count": len(daily_rows),
            "daily_amount": sum((row["amount_decimal"] for row in daily_rows), Decimal("0")),
            "annual_source_rows": source_rows_text(annual_rows),
            "annual_row_count": len(annual_rows),
            "annual_amount": sum((row["amount_decimal"] for row in annual_rows), Decimal("0")),
        })

    changed = []
    for contract_number, year_candidates in year_end_view.items():
        latest_candidates = full_latest_view[contract_number]
        if {row["source_row"] for row in year_candidates} != {row["source_row"] for row in latest_candidates}:
            payment_rows_for_contract = annual_by_contract.get(contract_number, [])
            annual_amount = sum((row["amount_decimal"] for row in payment_rows_for_contract), Decimal("0"))
            changed.append({
                "comparison_type": "年末既有合同被未来版本替换",
                "contract_number": contract_number,
                "annual_payment_amount": annual_amount,
                "annual_payment_source_rows": source_rows_text(payment_rows_for_contract),
                "year_end_contract_rows": source_rows_text(year_candidates),
                "year_end_approval": " | ".join(sorted({row["approval_date"] for row in year_candidates})),
                "year_end_revision": " | ".join(sorted({row["revision_number"] for row in year_candidates}, key=int)),
                "year_end_award": " | ".join(sorted({decimal_text(row["award_decimal"]) if row["award_decimal"] is not None else "缺失" for row in year_candidates})),
                "latest_contract_rows": source_rows_text(latest_candidates),
                "latest_approval": " | ".join(sorted({row["approval_date"] for row in latest_candidates})),
                "latest_revision": " | ".join(sorted({row["revision_number"] for row in latest_candidates}, key=int)),
                "latest_award": " | ".join(sorted({decimal_text(row["award_decimal"]) if row["award_decimal"] is not None else "缺失" for row in latest_candidates})),
            })
    for contract_number in set(full_latest_view) - set(year_end_view):
        latest_candidates = full_latest_view[contract_number]
        payment_rows_for_contract = annual_by_contract.get(contract_number, [])
        annual_amount = sum((row["amount_decimal"] for row in payment_rows_for_contract), Decimal("0"))
        changed.append({
            "comparison_type": "全快照最新法引入年末不存在合同",
            "contract_number": contract_number,
            "annual_payment_amount": annual_amount,
            "annual_payment_source_rows": source_rows_text(payment_rows_for_contract),
            "year_end_contract_rows": "",
            "year_end_approval": "",
            "year_end_revision": "",
            "year_end_award": "",
            "latest_contract_rows": source_rows_text(latest_candidates),
            "latest_approval": " | ".join(sorted({row["approval_date"] for row in latest_candidates})),
            "latest_revision": " | ".join(sorted({row["revision_number"] for row in latest_candidates}, key=int)),
            "latest_award": " | ".join(sorted({decimal_text(row["award_decimal"]) if row["award_decimal"] is not None else "缺失" for row in latest_candidates})),
        })
    changed.sort(key=lambda item: (-abs(item["annual_payment_amount"]), item["contract_number"]))

    unmatched = []
    high_ratio = []
    award_exceptions = []
    for contract_number, rows in annual_by_contract.items():
        annual_amount = sum((row["amount_decimal"] for row in rows), Decimal("0"))
        status = payment_status(rows[0], year_end_view)
        if status in {"未匹配", "版本冲突", "直接付款/空号"}:
            candidates = year_end_view.get(contract_number, [])
            if status == "未匹配":
                gap = "截至2024-12-31的合同快照中无该合同号"
                next_step = "调取合同封面、截至年末的批准/修订记录及对应发票；核对合同号录入"
            elif status == "版本冲突":
                gap = "最高批准日期和修订号仍有内容冲突，不能唯一选版"
                next_step = "调取两条原始合同记录及采购系统版本审计轨迹，确认年末有效版本"
            else:
                gap = "付款标识为DV或合同号为空，不应按名称强配"
                next_step = "调取付款凭证、发票和采购依据，确认是否属于直接付款及其授权"
            unmatched.append({
                "contract_number": contract_number or "(空号)",
                "payment_source_rows": source_rows_text(rows),
                "payment_row_count": len(rows),
                "payment_vendor_names": vendors_text(rows),
                "annual_payment_amount": annual_amount,
                "match_status": status,
                "contract_source_rows": source_rows_text(candidates),
                "approval_date": " | ".join(row["approval_date"] for row in candidates),
                "revision_number": " | ".join(row["revision_number"] for row in candidates),
                "evidence_gap": gap,
                "next_step": next_step,
            })
        elif status == "唯一匹配":
            selected = year_end_view[contract_number][0]
            award = selected["award_decimal"]
            common = {
                "contract_number": contract_number,
                "payment_source_rows": source_rows_text(rows),
                "payment_row_count": len(rows),
                "payment_vendor_names": vendors_text(rows),
                "annual_payment_amount": annual_amount,
                "contract_source_row": selected["source_row"],
                "approval_date": selected["approval_date"],
                "revision_number": selected["revision_number"],
                "contract_vendor_name": selected["vendor_name"],
                "award_amount": award,
            }
            if award is not None and award > 0:
                high_ratio.append({
                    **common,
                    "payment_award_ratio": annual_amount / award,
                    "evidence_gap": "年度付款高于该时点所选版本的授予额；快照不含完整合同档案、发票或授予额字段释义",
                    "next_step": "调取截至年末全部批准修订、订单/发票和付款明细；确认授予额字段口径及是否跨期/跨版本",
                })
            else:
                category = "授予额缺失" if award is None else "授予额为零" if award == 0 else "授予额为负"
                award_exceptions.append({
                    **common,
                    "exception_type": category,
                    "evidence_gap": f"{category}，不能计算有意义的付款/授予额比率",
                    "next_step": "调取合同原件和修订历史，核对授予额字段含义、调整方向及累计授权金额",
                })
    unmatched.sort(key=lambda item: (-abs(item["annual_payment_amount"]), item["contract_number"]))
    high_ratio.sort(key=lambda item: (-item["payment_award_ratio"], item["contract_number"]))
    award_exceptions.sort(key=lambda item: (-abs(item["annual_payment_amount"]), item["contract_number"]))

    year_conflicts = []
    for contract_number, candidates in sorted(year_end_view.items()):
        if len(candidates) > 1:
            for row in candidates:
                year_conflicts.append({
                    "contract_number": contract_number,
                    "contract_source_row": row["source_row"],
                    "approval_date": row["approval_date"],
                    "revision_number": row["revision_number"],
                    "vendor_name": row["vendor_name"],
                    "award_amount": row["award_decimal"],
                })

    negative_payments = [{
        "source_row": row["source_row"], "grain": "逐日" if row in daily else "年度汇总",
        "voucher_number": row["voucher_number"], "check_date": row["check_date"],
        "contract_number": row["contract_number"], "vendor_name": row["vendor_name"],
        "amount": row["amount_decimal"],
    } for row in payments if row["amount_decimal"] < 0]

    full_contract_key = duplicate_stats(contracts, contract_fields)
    contract_business_key = duplicate_stats(contracts, ["purchase_order_contract_number", "revision_number", "approval_date"])
    daily_exact = duplicate_stats(daily, payment_fields)
    annual_exact = duplicate_stats(annual, payment_fields)
    daily_voucher = duplicate_stats(daily, ["voucher_number"])
    annual_candidate = duplicate_stats(annual, ["department_name", "contract_number", "vendor_name", "check_date"])

    scope_rows = []
    for grain, rows in [("逐日付款", daily), ("年度汇总", annual), ("无法解释日期", unexplained_dates)]:
        scope_rows.append({
            "grain": grain,
            "row_count": len(rows),
            "amount": sum((row["amount_decimal"] for row in rows), Decimal("0")),
            "negative_rows": sum(row["amount_decimal"] < 0 for row in rows),
            "negative_amount": sum((row["amount_decimal"] for row in rows if row["amount_decimal"] < 0), Decimal("0")),
            "dv_rows": sum(row["contract_number"] == "DV" for row in rows),
            "blank_contract_rows": sum(not row["contract_number"] for row in rows),
            "distinct_contract_markers": len({row["contract_number"] for row in rows}),
        })

    return {
        "daily": daily,
        "annual": annual,
        "unexplained_dates": unexplained_dates,
        "year_end_view": year_end_view,
        "full_latest_view": full_latest_view,
        "scope_rows": scope_rows,
        "daily_status": summarize_status(daily, year_end_view),
        "annual_status": summarize_status(annual, year_end_view),
        "overlap": overlap,
        "changed": changed,
        "unmatched_top20": unmatched[:20],
        "high_ratio_top20": high_ratio[:20],
        "award_exceptions": award_exceptions,
        "year_conflicts": year_conflicts,
        "negative_payments": negative_payments,
        "diagnostics": {
            "payments": {
                "record_grain": "check_date为MM/DD/YYYY的是逐日付款查询行；check_date为2024的是年度汇总行；两者分开报告",
                "stable_key": "CSV物理行号（表头为第1行）",
                "daily_candidate_key": "无可证实唯一业务键；voucher_number重复，完整字段也有重复",
                "annual_candidate_key": "无可证实唯一业务键；department_name+contract_number+vendor_name+check_date仍可能重复",
                "all_rows": len(payments), "daily_exact_duplicates": daily_exact,
                "annual_exact_duplicates": annual_exact, "daily_voucher_duplicates": daily_voucher,
                "annual_candidate_duplicates": annual_candidate,
            },
            "contracts": {
                "record_grain": "批准合同版本/修订记录；一个合同号可有多版本，版本数不等于合同数",
                "stable_key": "CSV物理行号（表头为第1行）",
                "candidate_business_key": "合同号+修订号+批准日期，但存在重复/内容冲突，不能单独作为唯一键",
                "all_rows": len(contracts), "distinct_contracts": len({row["purchase_order_contract_number"] for row in contracts}),
                "full_row_duplicates": full_contract_key, "business_key_duplicates": contract_business_key,
            },
        },
    }


def json_ready(value):
    if isinstance(value, Decimal):
        return decimal_text(value)
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dict):
        return {key: json_ready(item) for key, item in value.items()}
    if isinstance(value, list):
        return [json_ready(item) for item in value]
    return value


def result_payload(analysis):
    daily_total = sum((row["amount_decimal"] for row in analysis["daily"]), Decimal("0"))
    annual_total = sum((row["amount_decimal"] for row in analysis["annual"]), Decimal("0"))
    changed_with_payments = [item for item in analysis["changed"] if item["annual_payment_source_rows"]]
    changed_net = sum((item["annual_payment_amount"] for item in changed_with_payments), Decimal("0"))
    changed_abs = sum((abs(item["annual_payment_amount"]) for item in changed_with_payments), Decimal("0"))
    return {
        "metadata": {
            "cutoff": CUTOFF.isoformat(),
            "inputs": [PAYMENTS_FILE, CONTRACTS_FILE],
            "source_row_rule": "CSV表头为第1行，数据从第2行起",
            "limitations": "仅评价两份公开快照；不推断快照外合同、发票、银行流水或违法性",
        },
        "record_grain_and_keys": analysis["diagnostics"],
        "payment_scope": analysis["scope_rows"],
        "conservation": {
            "daily": {"before_join": daily_total, "after_join": sum((item["amount"] for item in analysis["daily_status"]), Decimal("0")), "difference": daily_total - sum((item["amount"] for item in analysis["daily_status"]), Decimal("0"))},
            "annual": {"before_join": annual_total, "after_join": sum((item["amount"] for item in analysis["annual_status"]), Decimal("0")), "difference": annual_total - sum((item["amount"] for item in analysis["annual_status"]), Decimal("0"))},
        },
        "join_summary": {"daily": analysis["daily_status"], "annual": analysis["annual_status"]},
        "cross_grain_overlap": analysis["overlap"],
        "contract_views": {
            "year_end_contract_count": len(analysis["year_end_view"]),
            "year_end_conflict_contract_count": len({item["contract_number"] for item in analysis["year_conflicts"]}),
            "full_latest_contract_count": len(analysis["full_latest_view"]),
            "different_contract_count_total": len(analysis["changed"]),
            "existing_year_end_contracts_replaced": sum(item["comparison_type"] == "年末既有合同被未来版本替换" for item in analysis["changed"]),
            "future_contracts_introduced": sum(item["comparison_type"] == "全快照最新法引入年末不存在合同" for item in analysis["changed"]),
            "different_contracts_with_annual_payments": len(changed_with_payments),
            "annual_payment_exposure_net": changed_net,
            "annual_payment_exposure_absolute": changed_abs,
            "top10_by_absolute_annual_payment": analysis["changed"][:10],
            "year_end_conflict_candidates": analysis["year_conflicts"],
        },
        "unmatched_follow_up_top20": analysis["unmatched_top20"],
        "high_ratio_follow_up_top20": analysis["high_ratio_top20"],
        "nonpositive_or_missing_award_contracts": analysis["award_exceptions"],
        "payment_records": [{
            "source_row": row["source_row"],
            "grain": grain,
            "voucher_number": row["voucher_number"],
            "check_date": row["check_date"],
            "department_name": row["department_name"],
            "contract_number": row["contract_number"],
            "vendor_name": row["vendor_name"],
            "amount": row["amount_decimal"],
            "negative_flag": row["amount_decimal"] < 0,
            "direct_or_blank_flag": row["contract_number"] in DIRECT_MARKERS,
            "join_status": payment_status(row, analysis["year_end_view"]),
        } for grain, rows in [("逐日付款", analysis["daily"]), ("年度汇总", analysis["annual"]), ("无法解释日期", analysis["unexplained_dates"])] for row in rows],
        "negative_payments": analysis["negative_payments"],
    }


def add_sheet(workbook, title, headers, rows, money_columns=(), ratio_columns=()):
    ws = workbook.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in money_columns:
        if column in headers:
            index = headers.index(column) + 1
            for cell in ws.iter_cols(min_col=index, max_col=index, min_row=2):
                cell[0].number_format = MONEY_FMT
    for column in ratio_columns:
        if column in headers:
            index = headers.index(column) + 1
            for cell in ws.iter_cols(min_col=index, max_col=index, min_row=2):
                cell[0].number_format = RATIO_FMT
    for index, header in enumerate(headers, 1):
        values = [str(ws.cell(row=row, column=index).value or "") for row in range(1, min(ws.max_row, 100) + 1)]
        width = min(max(max(map(len, values)) + 2, 12), 45)
        ws.column_dimensions[get_column_letter(index)].width = width
    if rows:
        table = Table(displayName=f"T{len(workbook.worksheets)}", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    return ws


def write_workbook(path, analysis):
    wb = Workbook()
    wb.remove(wb.active)
    scope = []
    for item in analysis["scope_rows"]:
        scope.append({"类别": "付款范围", "口径/状态": item["grain"], "行数": item["row_count"], "合同标识数": item["distinct_contract_markers"], "金额": item["amount"], "负数行": item["negative_rows"], "负数金额": item["negative_amount"], "DV行": item["dv_rows"], "空合同号行": item["blank_contract_rows"], "说明": "不同付款粒度不合并去重"})
    for grain, items in [("逐日付款", analysis["daily_status"]), ("年度汇总", analysis["annual_status"])]:
        before = sum((row["amount_decimal"] for row in analysis["daily"] if grain == "逐日付款"), Decimal("0")) if grain == "逐日付款" else sum((row["amount_decimal"] for row in analysis["annual"]), Decimal("0"))
        for item in items:
            scope.append({"类别": f"{grain}关联", "口径/状态": item["status"], "行数": item["row_count"], "合同标识数": item["contract_count"], "金额": item["amount"], "负数行": None, "负数金额": None, "DV行": None, "空合同号行": None, "说明": "按源付款行分类，不展开合同候选"})
        after = sum((item["amount"] for item in items), Decimal("0"))
        scope.append({"类别": f"{grain}守恒", "口径/状态": "关联前后差额", "行数": None, "合同标识数": None, "金额": before - after, "负数行": None, "负数金额": None, "DV行": None, "空合同号行": None, "说明": f"关联前={before}；关联后={after}"})
    d = analysis["diagnostics"]
    scope.extend([
        {"类别": "键与重复", "口径/状态": "付款稳定键", "说明": d["payments"]["stable_key"] + "；" + d["payments"]["daily_candidate_key"]},
        {"类别": "键与重复", "口径/状态": "逐日完整字段重复", "行数": d["payments"]["daily_exact_duplicates"]["duplicate_extra_rows"], "说明": str(d["payments"]["daily_exact_duplicates"])},
        {"类别": "键与重复", "口径/状态": "年度完整字段重复", "行数": d["payments"]["annual_exact_duplicates"]["duplicate_extra_rows"], "说明": str(d["payments"]["annual_exact_duplicates"])},
        {"类别": "键与重复", "口径/状态": "合同稳定键", "说明": d["contracts"]["stable_key"] + "；" + d["contracts"]["candidate_business_key"]},
        {"类别": "键与重复", "口径/状态": "合同候选业务键重复", "行数": d["contracts"]["business_key_duplicates"]["duplicate_extra_rows"], "说明": str(d["contracts"]["business_key_duplicates"])},
    ])
    add_sheet(wb, "口径与守恒", ["类别", "口径/状态", "行数", "合同标识数", "金额", "负数行", "负数金额", "DV行", "空合同号行", "说明"], scope, ["金额", "负数金额"])

    payment_detail = []
    for grain, rows in [("逐日付款", analysis["daily"]), ("年度汇总", analysis["annual"]), ("无法解释日期", analysis["unexplained_dates"])]:
        for row in rows:
            payment_detail.append({
                "付款原始行": row["source_row"], "粒度": grain, "凭证号": row["voucher_number"],
                "支票日期/年度标识": row["check_date"], "部门": row["department_name"],
                "合同号/直接付款标识": row["contract_number"], "供应商名": row["vendor_name"],
                "金额": row["amount_decimal"], "负数标识": "是" if row["amount_decimal"] < 0 else "否",
                "DV/空号标识": "是" if row["contract_number"] in DIRECT_MARKERS else "否",
                "年末合同关联状态": payment_status(row, analysis["year_end_view"]),
            })
    add_sheet(wb, "付款明细", ["付款原始行", "粒度", "凭证号", "支票日期/年度标识", "部门", "合同号/直接付款标识", "供应商名", "金额", "负数标识", "DV/空号标识", "年末合同关联状态"], payment_detail, ["金额"])

    version_rows = []
    for item in analysis["changed"]:
        version_rows.append({
            "偏差类型": item["comparison_type"], "合同号": item["contract_number"], "年度付款原始行": item["annual_payment_source_rows"], "年度付款净额": item["annual_payment_amount"],
            "年末合同原始行": item["year_end_contract_rows"], "年末批准日期": item["year_end_approval"], "年末修订号": item["year_end_revision"], "年末授予额": item["year_end_award"],
            "全快照最新原始行": item["latest_contract_rows"], "最新批准日期": item["latest_approval"], "最新修订号": item["latest_revision"], "最新授予额": item["latest_award"],
            "说明": "全快照最新法与年末视图不同；未来版本不代表年末已获批，年末空白表示该合同截至年末尚无批准记录",
        })
    add_sheet(wb, "版本对照", ["偏差类型", "合同号", "年度付款原始行", "年度付款净额", "年末合同原始行", "年末批准日期", "年末修订号", "年末授予额", "全快照最新原始行", "最新批准日期", "最新修订号", "最新授予额", "说明"], version_rows, ["年度付款净额"])

    unmatched_rows = [{
        "合同号": x["contract_number"], "付款原始行": x["payment_source_rows"], "付款行数": x["payment_row_count"], "付款供应商名": x["payment_vendor_names"], "年度付款净额": x["annual_payment_amount"], "匹配状态": x["match_status"],
        "合同原始行": x["contract_source_rows"], "批准日期": x["approval_date"], "修订号": x["revision_number"], "证据缺口": x["evidence_gap"], "具体下一步": x["next_step"]
    } for x in analysis["unmatched_top20"]]
    add_sheet(wb, "未匹配追查", ["合同号", "付款原始行", "付款行数", "付款供应商名", "年度付款净额", "匹配状态", "合同原始行", "批准日期", "修订号", "证据缺口", "具体下一步"], unmatched_rows, ["年度付款净额"])

    ratio_rows = [{
        "合同号": x["contract_number"], "付款原始行": x["payment_source_rows"], "付款行数": x["payment_row_count"], "付款供应商名": x["payment_vendor_names"], "年度付款净额": x["annual_payment_amount"],
        "合同原始行": x["contract_source_row"], "批准日期": x["approval_date"], "修订号": x["revision_number"], "合同供应商名": x["contract_vendor_name"], "授予额": x["award_amount"], "付款/授予额": x["payment_award_ratio"], "证据缺口": x["evidence_gap"], "具体下一步": x["next_step"]
    } for x in analysis["high_ratio_top20"]]
    add_sheet(wb, "高比率追查", ["合同号", "付款原始行", "付款行数", "付款供应商名", "年度付款净额", "合同原始行", "批准日期", "修订号", "合同供应商名", "授予额", "付款/授予额", "证据缺口", "具体下一步"], ratio_rows, ["年度付款净额", "授予额"], ["付款/授予额"])

    exceptions = []
    for x in analysis["award_exceptions"]:
        exceptions.append({"例外类型": x["exception_type"], "合同号": x["contract_number"], "付款原始行": x["payment_source_rows"], "付款日期/口径": "2024", "付款供应商名": x["payment_vendor_names"], "付款金额": x["annual_payment_amount"], "合同原始行": x["contract_source_row"], "批准日期": x["approval_date"], "修订号": x["revision_number"], "授予额": x["award_amount"], "证据缺口": x["evidence_gap"], "具体下一步": x["next_step"]})
    for x in analysis["negative_payments"]:
        exceptions.append({"例外类型": "负数付款/冲销（保留）", "合同号": x["contract_number"], "付款原始行": x["source_row"], "付款日期/口径": x["check_date"], "付款供应商名": x["vendor_name"], "付款金额": x["amount"], "合同原始行": None, "批准日期": None, "修订号": None, "授予额": None, "证据缺口": "快照不说明冲销对应的原付款", "具体下一步": "调取付款凭证和总账分录，核对冲销关联及会计期间"})
    for x in analysis["year_conflicts"]:
        exceptions.append({"例外类型": "年末版本冲突候选", "合同号": x["contract_number"], "付款原始行": None, "付款日期/口径": None, "付款供应商名": None, "付款金额": None, "合同原始行": x["contract_source_row"], "批准日期": x["approval_date"], "修订号": x["revision_number"], "授予额": x["award_amount"], "证据缺口": "最高批准日期和修订号相同但内容不一致", "具体下一步": "调取采购系统版本审计轨迹和合同原件，确认有效内容"})
    for x in analysis["overlap"]:
        exceptions.append({"例外类型": "跨粒度合同交集", "合同号": x["contract_number"], "付款原始行": f"逐日:{x['daily_source_rows']}；年度:{x['annual_source_rows']}", "付款日期/口径": "逐日与年度", "付款供应商名": None, "付款金额": x["daily_amount"] + x["annual_amount"], "合同原始行": None, "批准日期": None, "修订号": None, "授予额": None, "证据缺口": f"逐日净额{x['daily_amount']}；年度净额{x['annual_amount']}；无证据证明互斥", "具体下一步": "调取查询定义和底层交易标识，确认年度行是否汇总包含逐日行；在此之前不得相加为真实总支出"})
    add_sheet(wb, "例外", ["例外类型", "合同号", "付款原始行", "付款日期/口径", "付款供应商名", "付款金额", "合同原始行", "批准日期", "修订号", "授予额", "证据缺口", "具体下一步"], exceptions, ["付款金额", "授予额"])
    wb.save(path)


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.style.font.name = "Arial"
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def write_memo(path, analysis):
    daily_total = sum((row["amount_decimal"] for row in analysis["daily"]), Decimal("0"))
    annual_total = sum((row["amount_decimal"] for row in analysis["annual"]), Decimal("0"))
    changed_paid = [item for item in analysis["changed"] if item["annual_payment_source_rows"]]
    changed_net = sum((item["annual_payment_amount"] for item in changed_paid), Decimal("0"))
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    title = doc.add_heading("芝加哥2024年付款复核备忘录", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("研究截点：2024-12-31｜仅基于两份公开快照").bold = True

    daily_direct = next(item for item in analysis["daily_status"] if item["status"] == "直接付款/空号")
    annual_unique = next(item for item in analysis["annual_status"] if item["status"] == "唯一匹配")
    annual_unmatched = next(item for item in analysis["annual_status"] if item["status"] == "未匹配")
    add_heading(doc, "结论先行", 1)
    add_bullet(doc, f"付款快照混有两种不可直接相加的粒度：{len(analysis['daily']):,}条具体日期行，净额${daily_total:,.2f}；{len(analysis['annual']):,}条check_date='2024'年度汇总行，净额${annual_total:,.2f}。{len(analysis['overlap'])}个合同号同时出现在两类中；没有底层交易标识证明互斥，因此本复核不报告消除重叠后的“真实总支出”。")
    add_bullet(doc, f"截至年末可选出{len(analysis['year_end_view']):,}个合同号；{len({item['contract_number'] for item in analysis['year_conflicts']}):,}个合同号在最高批准日期和修订号下仍有内容冲突候选，未擅选。年度付款中，{annual_unique['row_count']:,}行/{annual_unique['contract_count']:,}个合同号/净额${annual_unique['amount']:,.2f}唯一匹配，{annual_unmatched['row_count']:,}行/{annual_unmatched['contract_count']:,}个合同号/净额${annual_unmatched['amount']:,.2f}未匹配。")
    replaced_count = sum(item["comparison_type"] == "年末既有合同被未来版本替换" for item in analysis["changed"])
    introduced_count = sum(item["comparison_type"] == "全快照最新法引入年末不存在合同" for item in analysis["changed"])
    add_bullet(doc, f"若直接采用全快照最新版本，与年末视图不同的合同共{len(analysis['changed']):,}个：{replaced_count:,}个年末既有合同被未来版本替换，另引入{introduced_count:,}个截至年末尚不存在的合同；其中{len(changed_paid):,}个带有年度付款，付款敞口净额${changed_net:,.2f}。这是时点偏差，不是付款错误。")
    add_bullet(doc, "优先补档分两组：先查年度未唯一匹配且绝对净额最大的合同；再查唯一匹配但付款/正授予额比率最高的合同。高比率只表示快照字段之间需要解释，不等于违法、超付或损失。")

    add_heading(doc, "已核实事实", 1)
    add_bullet(doc, "付款CSV以物理行号作为稳定来源键。voucher_number及完整字段均存在重复，不能假定凭证号唯一；重复记录未删除。合同CSV同样以物理行号为稳定键，合同号+修订号+批准日期有9个重复组，完整行无重复。")
    add_bullet(doc, "付款中保留两条负数：原始行8720为逐日DV付款-$500.00；原始行39009为合同192984年度汇总-$4,051.61。负数被视为冲销/调整线索并纳入净额。")
    add_bullet(doc, "年末版本规则严格为：先排除批准日期晚于2024-12-31的记录，再按合同号选批准日期最新、其后数值修订号最大；最高顺序内容冲突全部保留。授予额是所选版本字段，不是累计付款。")
    add_bullet(doc, f"逐日付款中直接付款/空号共{daily_direct['row_count']:,}行，净额${daily_direct['amount']:,.2f}，独立归类且未按供应商名称强配。两个粒度关联后金额差额均为$0.00。")

    add_heading(doc, "优先追查与完整证据链示范", 1)
    unmatched_example = analysis["unmatched_top20"][0]
    ratio_example = analysis["high_ratio_top20"][0]
    changed_example = analysis["changed"][0]
    conflict_contract = analysis["year_conflicts"][0]["contract_number"] if analysis["year_conflicts"] else None
    examples = [
        (unmatched_example["contract_number"], "未匹配", f"付款原始行{unmatched_example['payment_source_rows']}记录年度净额${unmatched_example['annual_payment_amount']:,.2f}、供应商{unmatched_example['payment_vendor_names']}；年末合同视图无该合同号。下一步取得合同封面、年末前批准/修订、发票及付款凭证，并核对付款合同号录入。现有证据只能说明快照间未匹配。"),
        (ratio_example["contract_number"], "高比率", f"付款原始行{ratio_example['payment_source_rows']}记录年度净额${ratio_example['annual_payment_amount']:,.2f}；年末所选合同原始行{ratio_example['contract_source_row']}，批准日{ratio_example['approval_date'][:10]}、修订{ratio_example['revision_number']}、授予额${ratio_example['award_amount']:,.2f}，付款与合同供应商文本分别为“{ratio_example['payment_vendor_names']}”和“{ratio_example['contract_vendor_name']}”，比率{ratio_example['payment_award_ratio']:,.2f}倍。下一步核对合同原件、授予额字段口径、全部修订及发票；不能仅据此认定超付。"),
        (changed_example["contract_number"], "未来版本偏差", f"付款原始行为{changed_example['annual_payment_source_rows'] or '无年度付款行'}，年度净额${changed_example['annual_payment_amount']:,.2f}；年末版本为合同原始行{changed_example['year_end_contract_rows']}（{changed_example['year_end_approval'][:10]}、修订{changed_example['year_end_revision']}、授予额{changed_example['year_end_award']}），全快照最新变为原始行{changed_example['latest_contract_rows']}（{changed_example['latest_approval'][:10]}、修订{changed_example['latest_revision']}、授予额{changed_example['latest_award']}）。采用最新法会把未来批准版本倒灌至年末。"),
    ]
    if conflict_contract:
        candidates = [item for item in analysis["year_conflicts"] if item["contract_number"] == conflict_contract]
        details = "；".join(f"原始行{item['contract_source_row']}，{item['approval_date'][:10]}、修订{item['revision_number']}、供应商“{item['vendor_name']}”、授予额${item['award_amount']:,.2f}" for item in candidates)
        examples.append((conflict_contract, "版本冲突", f"{details}。最高批准日期和修订号相同但内容不一致，年末规则不能唯一选择；应调取采购系统版本审计轨迹和合同原件。"))
    for number, label, detail in examples:
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"合同{number}（{label}）：").bold = True
        paragraph.add_run(detail)

    add_heading(doc, "需要补档的异常", 1)
    top_unmatched = analysis["unmatched_top20"][:2]
    top_unmatched_text = "和".join(f"合同{item['contract_number']}（${item['annual_payment_amount']:,.2f}）" for item in top_unmatched)
    add_bullet(doc, f"“未匹配追查”前20按年度付款绝对净额降序，首两项为{top_unmatched_text}。应取得合同号主数据、合同封面、年末前批准修订与发票链。")
    add_bullet(doc, "“高比率追查”前20按付款/授予额降序；应确认授予额字段是否代表净变更、单一修订或其他口径，以及付款是否跨期、跨订单或跨版本。名称差异仅记录为证据问题，不用于认定实体同一。")
    zero_count = sum(item["exception_type"] == "授予额为零" for item in analysis["award_exceptions"])
    negative_count = sum(item["exception_type"] == "授予额为负" for item in analysis["award_exceptions"])
    missing_count = sum(item["exception_type"] == "授予额缺失" for item in analysis["award_exceptions"])
    add_bullet(doc, f"另有{zero_count}个付款合同的所选授予额为零、{negative_count}个为负、{missing_count}个缺失；这些项目在“例外”中单列，不制造无穷比率。")

    add_heading(doc, "不能下的结论", 1)
    add_bullet(doc, "不能把付款高于所选版本授予额直接写成违法、超付、浪费或欺诈；快照不含完整合同档案、发票、采购订单、总账或银行流水。")
    add_bullet(doc, "不能将全快照最新合同视为2024年末已获批版本；合同表含2025—2026年批准记录。")
    add_bullet(doc, "不能把逐日净额与年度汇总净额简单相加，也不能在无交易级映射时消除重叠后声称得到真实总支出。")
    add_bullet(doc, "不能依据供应商名称相似或不同强制匹配、拆分或合并实体；DV/空号必须靠付款凭证和采购依据补证。")

    add_heading(doc, "范围与复核路径", 1)
    doc.add_paragraph("本包仅复核Chicago_Payments_2024.csv与Chicago_Contracts_2019plus.csv的下载快照。audit.xlsx提供口径守恒、版本对照、两张前20清单和全部例外；results.json保留结构化数值和来源行号；analysis.py可离线重建三份产物。前20是取证顺序，不是违规排序。")
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    doc.save(path)


def main():
    parser = argparse.ArgumentParser(description="重建芝加哥2024付款复核交付物")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    if not args.output.exists() or not args.output.is_dir():
        raise ValueError("输出目录必须预先存在")
    payment_fields, payments = payment_rows(args.input / PAYMENTS_FILE)
    contract_fields, contracts = contract_rows(args.input / CONTRACTS_FILE)
    analysis = build_analysis(payment_fields, payments, contract_fields, contracts)
    payload = result_payload(analysis)
    with (args.output / "results.json").open("w", encoding="utf-8") as handle:
        json.dump(json_ready(payload), handle, ensure_ascii=False, indent=2)
    write_workbook(args.output / "audit.xlsx", analysis)
    write_memo(args.output / "memo.docx", analysis)


if __name__ == "__main__":
    main()
