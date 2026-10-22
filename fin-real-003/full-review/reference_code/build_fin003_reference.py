"""Compute a private, reproducible reference candidate from the frozen originals.

This is an author candidate, not an independently approved Gold solution.
All authored tables are outputs; source files are never edited.
"""
from collections import Counter, defaultdict
import csv
from decimal import Decimal as D
import hashlib
import json
from pathlib import Path
from statistics import median

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent/'redesign/FIN-REAL-003'
METRICS = ['roa_asset_proxy', 'nonaccrual_assets_pct', 'nco_loan_proxy',
           'rbc_asset_proxy', 'equity_asset_pct', 'loan_deposit_pct',
           'real_estate_loan_pct', 'nim_asset_proxy']
LIMITS = (
    '仅代表冻结导出的ME/NH/VT样本；当前机构表仅用作身份、业务特征和沿革核验，'
    '不以ACTIVE回删历史样本，也不使用机构表财务值覆盖季度数据。'
    'ROA/NIM已有年化；以期末资产加权仅为平均资产权重的近似。'
    'RBC以期末资产加权仅为缺少RWAJT时的近似，不是组合监管资本充足率。'
    'RBCRWAJ=0的记录保留原始值但标记有效性未确认，既不认定为真实零资本，也不无说明混入比率；资本代理仅对正值样本归一加权，需单列覆盖。'
    '净核销采用NTLNLS年初至今金额，按12/月数年化；期末贷款或资产代替平均分母仅为代理。'
    'NCLNLSR是非流动贷款率，不是净核销率，ELNATRY也不是净核销率。'
    '当前业务分类不证明历史业务始终不变；数学分解不是因果识别。'
    '同业偏离和内部触发条件不是监管评级、法定缺陷或要求调整准备的监管门槛。'
)


def read(name):
    with (ROOT/'input_files'/name).open(encoding='utf-8-sig', newline='') as f:
        return [dict(r, source_row=n) for n,r in enumerate(csv.DictReader(f),2)]


def num(r,k):
    return D(r[k])


def ratio(a,b):
    return a/b*100 if b else None


def build():
    raw=read('fdic_financials_VT_NH_ME_2022_2025.csv')
    inst=read('fdic_institutions_VT_NH_ME.csv')
    ix=defaultdict(list)
    for r in inst: ix[r['CERT']].append(r)
    dates=sorted({r['REPDTE'] for r in raw})
    bycert=defaultdict(list)
    for r in raw: bycert[r['CERT']].append(r)
    fixed={k for k,v in bycert.items() if len({r['REPDTE'] for r in v})==len(dates)}
    if any(n!=1 for n in Counter((r['CERT'],r['REPDTE']) for r in raw).values()):
        raise ValueError('Nonunique bank-quarter')
    if any(len(ix[k])!=1 for k in bycert): raise ValueError('Nonunique institution join')
    metadata={k:ix[k][0] for k in bycert}
    special={k for k,r in metadata.items() if 'Other Specialized' in r['SPECGRPN'] or 'Credit-card' in r['SPECGRPN']}
    conventional=fixed-special
    panel=[]
    for r in raw:
        annual=D(12)/D(int(r['REPDTE'][4:6]))
        a, loans=num(r,'ASSET'),num(r,'LNLSGR')
        panel.append(dict(cert=r['CERT'],name=r['NAMEFULL'],state=r['STALP'],date=r['REPDTE'],
            financial_row=r['source_row'],institution_row=metadata[r['CERT']]['source_row'],
            full_history=r['CERT'] in fixed,comparable=r['CERT'] in conventional,
            business=metadata[r['CERT']]['SPECGRPN'],bkclass=metadata[r['CERT']]['BKCLASS'],
            asset=a,loans=loans,dep=num(r,'DEP'),equity=num(r,'EQ'),na=num(r,'NAASSET'),
            nco_annual=num(r,'NTLNLS')*annual,re_loans=num(r,'LNRE'),net_loans=num(r,'LNLSNET'),
            roa_asset_proxy=num(r,'ROA'),nonaccrual_assets_pct=ratio(num(r,'NAASSET'),a),
            nco_loan_proxy=ratio(num(r,'NTLNLS')*annual,loans),
            nco_asset_proxy=ratio(num(r,'NTLNLS')*annual,a),
            rbc_raw=num(r,'RBCRWAJ'),rbc_validity='positive_reported' if num(r,'RBCRWAJ')>0 else 'zero_unresolved',
            rbc_asset_proxy=num(r,'RBCRWAJ') if num(r,'RBCRWAJ')>0 else None,equity_asset_pct=ratio(num(r,'EQ'),a),
            loan_deposit_pct=ratio(num(r,'LNLSNET'),num(r,'DEP')),
            real_estate_loan_pct=ratio(num(r,'LNRE'),loans),nim_asset_proxy=num(r,'NIMY')))

    def aggregate(rows):
        def total(k): return sum((r[k] for r in rows),D(0))
        def weighted(k,weight='asset'):
            eligible=[r for r in rows if r[k] is not None]
            return ratio(sum((r[k]*r[weight] for r in eligible),D(0)),sum((r[weight] for r in eligible),D(0)))/100 if eligible else None
        return dict(banks=len(rows),asset=total('asset'),dep=total('dep'),
            rbc_positive_banks=sum(r['rbc_asset_proxy'] is not None for r in rows),
            rbc_positive_asset=sum((r['asset'] for r in rows if r['rbc_asset_proxy'] is not None),D(0)),
            roa_asset_proxy=weighted('roa_asset_proxy'),rbc_asset_proxy=weighted('rbc_asset_proxy'),
            nonaccrual_assets_pct=ratio(total('na'),total('asset')),
            nco_loan_proxy=ratio(total('nco_annual'),total('loans')),
            nco_asset_proxy=ratio(total('nco_annual'),total('asset')),
            equity_asset_pct=ratio(total('equity'),total('asset')),
            loan_deposit_pct=ratio(total('net_loans'),total('dep')),
            real_estate_loan_pct=ratio(total('re_loans'),total('loans')),
            nim_asset_proxy=weighted('nim_asset_proxy'))

    aggregates=[]
    for date in dates:
        for state in ['ALL','ME','NH','VT']:
            for scope,ids in [('all',set(bycert)),('continuous',fixed),('continuous_conventional',conventional)]:
                rows=[r for r in panel if r['date']==date and r['cert'] in ids and (state=='ALL' or r['state']==state)]
                if rows: aggregates.append(dict(date=date,state=state,scope=scope,**aggregate(rows)))
    endpoints=[{r['cert']:r for r in panel if r['date']==q} for q in (dates[0],dates[-1])]
    common=set(endpoints[0])&set(endpoints[1]); decomp=[]; contributions=[]
    for metric,weight in [('roa_asset_proxy','asset'),('nonaccrual_assets_pct','asset'),
                          ('nco_loan_proxy','loans'),('rbc_asset_proxy','asset')]:
        scoped=[{k:r for k,r in e.items() if metric!='rbc_asset_proxy' or r[metric] is not None} for e in endpoints]
        common=set(scoped[0])&set(scoped[1])
        den=[sum((r[weight] for r in e.values()),D(0)) for e in scoped]
        w=[{k:r[weight]/d for k,r in e.items()} for e,d in zip(scoped,den)]
        def rate(j,k):
            value=endpoints[j][k][metric]
            if value is None:
                if w[j][k]!=0: raise ValueError('Missing positive-exposure metric')
                # Zero loan exposure contributes zero, without claiming a bank ratio exists.
                return D(0)
            return value
        terms=[]
        for k in sorted(set(scoped[0])|set(scoped[1])):
            within=reweight=entry=exit_=D(0)
            if k in common:
                r0,r1=rate(0,k),rate(1,k)
                within=(w[0][k]+w[1][k])/2*(r1-r0)
                reweight=(r0+r1)/2*(w[1][k]-w[0][k])
            elif k in scoped[1]: entry=w[1][k]*rate(1,k)
            else: exit_=-w[0][k]*rate(0,k)
            terms.append(dict(metric=metric,cert=k,within_pp=within,reweight_pp=reweight,entry_pp=entry,exit_pp=exit_))
        values=[sum((w[j][k]*rate(j,k) for k in e),D(0)) for j,e in enumerate(scoped)]
        sums={key:sum((t[key] for t in terms),D(0)) for key in ['within_pp','reweight_pp','entry_pp','exit_pp']}
        residual=values[1]-values[0]-sum(sums.values(),D(0))
        assert abs(residual)<D('1e-20')
        decomp.append(dict(metric=metric,weight=weight,start=values[0],end=values[1],change_pp=values[1]-values[0],**sums,residual=residual,
            scope='positive_RBC_observations_only; entry/exit includes validity coverage changes' if metric=='rbc_asset_proxy' else 'all financial observations'))
        contributions.extend(terms)

    events=[]
    for k,r in metadata.items():
        if r['NEWCERT'] in ('','0'): continue
        end=r['ENDEFYMD']; eventdate=end[6:]+end[:2]+end[3:5]
        own=sorted(bycert[k],key=lambda x:x['REPDTE'])
        successor=sorted(bycert.get(r['NEWCERT'],[]),key=lambda x:x['REPDTE'])
        before=[s for s in successor if s['REPDTE']<eventdate]
        after=[s for s in successor if s['REPDTE']>=eventdate]
        events.append(dict(cert=k,name=r['NAME'],successor=r['NEWCERT'],end_date=end,
            in_period=dates[0]<=eventdate<=dates[-1],last_financial_date=own[-1]['REPDTE'],last_asset=own[-1]['ASSET'],
            successor_before_date=before[-1]['REPDTE'] if before else None,
            successor_before_asset=before[-1]['ASSET'] if before else None,
            successor_after_date=after[0]['REPDTE'] if after else None,
            successor_after_asset=after[0]['ASSET'] if after else None,
            treatment='身份沿革支持结构性解释；前后金额不等于精确交易贡献，不拼造历史合并报表。'))
    names=defaultdict(set)
    for r in panel: names[r['name']].add(r['cert'])
    collisions=[dict(name=k,certs=sorted(v)) for k,v in names.items() if len(v)>1]
    coverage=[dict(cert=k,name=metadata[k]['NAME'],quarters=len(v),missing=sorted(set(dates)-{r['REPDTE'] for r in v}),
        comparable=k in conventional,business=metadata[k]['SPECGRPN'],active_current=metadata[k]['ACTIVE'],
        rationale='全期覆盖且非当前专营分类；分类有时点限制' if k in conventional else '单列覆盖不连续或特殊业务；保留于全样本') for k,v in bycert.items()]
    last=endpoints[1]; previous={r['cert']:r for r in panel if r['date']=='20240630'}
    eligible=[r for r in last.values() if r['cert'] in conventional]
    ranking=[]
    for r in eligible:
        signals={}
        for m,high in [('nonaccrual_assets_pct',True),('nco_loan_proxy',True),('roa_asset_proxy',False),('equity_asset_pct',False),('loan_deposit_pct',True)]:
            valid=[x[m] for x in eligible if x[m] is not None]
            signals[m]=D(sum(v<=r[m] if high else v>=r[m] for v in valid))/len(valid)
        before=previous.get(r['cert'])
        ranking.append(dict(cert=r['cert'],name=r['name'],score=sum(signals.values())/len(signals),
            **{m:r[m] for m in METRICS},signals=signals,
            yoy_na_pp=r['nonaccrual_assets_pct']-before['nonaccrual_assets_pct'] if before else None,
            yoy_roa_pp=r['roa_asset_proxy']-before['roa_asset_proxy'] if before else None))
    ranking.sort(key=lambda r:(-r['score'],int(r['cert'])))
    return dict(counts=dict(financial_rows=len(raw),quarters=len(dates),banks=len(bycert),matched_rows=len(raw),
                           continuous=len(fixed),conventional=len(conventional)),
        dates=dates,identity_collisions=collisions,coverage=coverage,panel=panel,aggregates=aggregates,
        decomposition=decomp,contributions=contributions,events=events,ranking=ranking,
        selected=ranking[:5],limitations=LIMITS,
        selection_method='连续可比机构内盈利、非应计、净核销、账面权益/资产、贷存比五维等权风险方向经验百分位，选前五用于有限资源初筛；RBC零值待核验不参与排名。非唯一正确名单、非监管评级。对排名边界与结构性事件另行复核。')


def write(result):
    dest=ROOT/'reference_output'; dest.mkdir(exist_ok=True)
    (dest/'results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str))
    wb=Workbook(); wb.remove(wb.active)
    sheets={k:result[k] for k in ['coverage','panel','aggregates','decomposition','contributions','events','ranking','identity_collisions']}
    sheets['readme']=[dict(item='用途',text='作者参考候选，未获独立Gold批准'),dict(item='限制',text=LIMITS),dict(item='筛选',text=result['selection_method'])]
    sheets['counts']=[result['counts']]
    for name,rows in sheets.items():
        ws=wb.create_sheet(name); keys=list(dict.fromkeys(k for row in rows for k in row))
        ws.append(keys)
        for row in rows:
            ws.append([float(row[k]) if isinstance(row.get(k),D) else json.dumps(row[k],ensure_ascii=False,default=str) if isinstance(row.get(k),(dict,list)) else row.get(k) for k in keys])
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='285C50')
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(55,max(14,max(len(str(c.value or '')) for c in col[:30])+2))
    wb.save(dest/'regional_bank_risk_model.xlsx')
    doc=Document(); doc.add_heading('经营表现诊断：作者参考候选',0)
    doc.add_paragraph('2022-03-31至2025-06-30。尚未独立审核，以下为一种可复算分析方案，不是唯一标准答案。')
    doc.add_heading('执行摘要',1)
    for row in result['decomposition']:
        doc.add_paragraph(f"{row['metric']}：{row['start']:.4f}% → {row['end']:.4f}%，变化{row['change_pp']:.4f}个百分点；自身表现{row['within_pp']:.4f}，权重{row['reweight_pp']:.4f}，进入{row['entry_pp']:.4f}，退出{row['exit_pp']:.4f}。")
    doc.add_paragraph(LIMITS)
    doc.add_heading('主要驱动与竞争性解释',1)
    for metric in ['roa_asset_proxy','nonaccrual_assets_pct','nco_loan_proxy','rbc_asset_proxy']:
        terms=sorted([r for r in result['contributions'] if r['metric']==metric],key=lambda r:abs(r['within_pp']),reverse=True)[:3]
        doc.add_paragraph(metric+'：持续机构自身变化绝对贡献前三为 '+ '；'.join(f"CERT {r['cert']}：{r['within_pp']:.4f}个百分点" for r in terms)+'。这是机械贡献，不能据此认定经营根因。')
    for e in result['events']:
        doc.add_paragraph(f"{e['name']}({e['cert']}) → {e['successor']}，结束日期{e['end_date']}，期内={e['in_period']}。退出前资产{e['last_asset']}；继承者前后资产{e['successor_before_asset']} / {e['successor_after_asset']}。{e['treatment']}")
    doc.add_paragraph('经营性解释是持续机构信用迁移或盈利压力；结构性解释是承继、进入退出及资产权重改变。分解和连续样本可区分机械贡献，却缺少借款人迁移、核销回收明细及交易会计桥接，无法证实贷款质量变化由何原因造成。')
    doc.add_heading('可比范围与行动取舍',1)
    ends=[r for r in result['aggregates'] if r['state']=='ALL' and r['date'] in (result['dates'][0],result['dates'][-1])]
    for r in ends: doc.add_paragraph(f"{r['date']} {r['scope']}：{r['banks']}家，ROA代理{r['roa_asset_proxy']:.4f}%，非应计资产率{r['nonaccrual_assets_pct']:.4f}%，资本代理{r['rbc_asset_proxy']:.4f}%。")
    doc.add_paragraph(result['selection_method'])
    doc.add_paragraph('选择分层问询而非全面相同深度复核：先复核五维信号较集中的机构，同时单列承继事件桥接；特殊业务机构不直接放入传统银行排名。边界银行不等于无风险，若差异化资料推翻当前信号、同业组改变或资本真实分母改变结论，应重排并扩大范围。')
    brief=Document(); brief.add_heading('优先跟进简报：作者参考候选',0)
    brief.add_paragraph(result['selection_method']); brief.add_paragraph(LIMITS)
    for n,r in enumerate(result['selected'],1):
        text=(f"{r['name']}，CERT {r['cert']}，内部初筛第{n}，综合方向分位{r['score']:.3f}。"
              f"ROA {r['roa_asset_proxy']:.3f}%，非应计资产率{r['nonaccrual_assets_pct']:.3f}%，"
              f"净核销贷款代理{r['nco_loan_proxy']:.3f}%，RBC银行比率{r['rbc_asset_proxy'] if r['rbc_asset_proxy'] is not None else '零值有效性待核验'}，"
              f"贷存比{r['loan_deposit_pct']:.2f}%，房地产贷款占比{r['real_estate_loan_pct']:.2f}%。"
              f"同比非应计变化{r['yoy_na_pp']:.3f}个百分点，ROA变化{r['yoy_roa_pp']:.3f}个百分点。")
        dominant=max(r['signals'],key=r['signals'].get)
        requests={
            'nonaccrual_assets_pct':'先问非应计增加是新增逾期、重组、并购接入还是回收变慢；请求逐笔迁移、担保与催收记录，追溯最高贡献组合。',
            'nco_loan_proxy':'先问核销时点、回收和资产出售是否造成年度可比差异；请求YTD核销及回收桥接、贷款平均余额，再验证年化代理。',
            'roa_asset_proxy':'先问利润变化来自利差、费用、一次性项目还是准备；请求月度损益、平均资产及净息差桥接，检验经常性盈利。',
            'equity_asset_pct':'先问账面权益承受力与监管资本是否一致、变化来自利润留存还是资产增长；请求RBC/RWAJT、资本行动与压力情景，确认监管资本数据有效性。',
            'loan_deposit_pct':'先问贷款增长和存款变化是否由客户集中或短期资金替代驱动；请求存款集中度、期限和融资成本、现金流压力及备用流动性。'}
        action=requests[dominant]
        for d in [doc,brief]:
            d.add_heading(f"{n}. {r['name']} ({r['cert']})",2); d.add_paragraph(text); d.add_paragraph(action)
            d.add_paragraph('补充请求贷款分行业/地区及关联集中度、逾期向非应计迁移、资本与资金承受力；若机构没有相关业务，应书面说明并缩减请求。由风险复核负责人将每项证据与最强信号和竞争解释对应，避免把资料存在本身当作根因已验证。')
    for d in [doc,brief]:
        d.add_heading('九十天安排与更新',1)
        for s in ['第1–15天：财务分析负责人核对身份、口径与原始快照，风险负责人向优先机构及承继事件相关机构发差异化资料请求；产出确认单。',
                  '第16–30天：信用、资金及资本条线复核迁移、回收、平均分母和交易桥接；产出可确认/可否定/待补证的解释清单，资料不足时不下定论。',
                  '第31–60天：管理层审议复核结果和替代同业分组敏感性，决定调整跟进顺序、扩大或缩减范围；产出责任方向明确的行动台账。',
                  '第61–90天：财务分析负责人更新季度数据，风险条线按月追踪入选银行信用迁移、资金和资本信号；新证据修正代理或承继解释时重算排序。']:
            d.add_paragraph(s)
        d.add_paragraph('逐家每月复核其最强信号及对应原始明细，组合按季度复算四项指标和贡献；当同业偏离持续、信用与资本/资金信号同时恶化或关键桥接长期无法解释时升级管理层复核；资料证实一次性结构变化且持续信号缓解时可降级。上述是拟议内部管理条件，不是既有政策或监管门槛。')
        d.add_paragraph('依据OCC Rating Credit Risk关于组合复杂度、根因及监测的原则，以及FDIC RMS Loans关于贷款质量和集中风险复核的原则，优先取得可验证资料。未获得逐笔证据，不认定正式监管评级，不按同业中位数直接要求调整信用损失准备。具体页码引用尚待独立审阅确认。')
    doc.save(dest/'operating_performance_diagnosis.docx'); brief.save(dest/'priority_follow_up_brief.docx')
    book=load_workbook(dest/'regional_bank_risk_model.xlsx',read_only=True)
    validation=dict(author_candidate=True,gold_approved=False,expert_review=False,
        workbook_sheets=len(book.sheetnames),panel_rows=book['panel'].max_row-1,
        decomposition_residuals=[str(r['residual']) for r in result['decomposition']],
        files={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in dest.iterdir() if p.is_file()},
        remaining=['独立数值与业务审核','Office实际渲染检查','rubric逐项证据与合理替代解审查'])
    book.close(); (dest/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2))
    print(json.dumps(dict(counts=result['counts'],selected=[r['cert'] for r in result['selected']],gold_approved=False)))


if __name__=='__main__': write(build())
