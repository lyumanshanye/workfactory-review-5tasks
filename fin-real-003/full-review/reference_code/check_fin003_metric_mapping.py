"""Record a specific, independently recomputed mapping error without assigning a score."""
import csv
from decimal import Decimal as D
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
TASK=ROOT/'redesign/FIN-REAL-003'
RUN=ROOT/'harness_runs/fin003-gpt-01'


def main():
    with (TASK/'input_files/fdic_financials_VT_NH_ME_2022_2025.csv').open() as f:
        rows=[r for r in csv.DictReader(f) if r['REPDTE']=='20250630']
    loans=sum(D(r['LNLSGR']) for r in rows)
    noncurrent=sum(D(r['NCLNLSR'])*D(r['LNLSGR']) for r in rows)/loans
    nco_proxy=sum(D(r['NTLNLS'])*2 for r in rows)/loans*100
    wb=load_workbook(RUN/'output/regional_bank_risk_model.xlsx',data_only=True,read_only=True)
    ws=wb['Executive Dashboard']; label=ws['A8'].value; candidate=D(str(ws['A9'].value)); wb.close()
    assert label=='Net charge-off rate (%)'
    assert abs(candidate-noncurrent)<D('1e-12')
    assert abs(candidate-nco_proxy)>D('.4')
    result=dict(finding='Noncurrent-loan ratio mislabeled as net charge-off rate',
        file='regional_bank_risk_model.xlsx',label_cell='Executive Dashboard!A8',value_cell='Executive Dashboard!A9',
        candidate_pct=str(candidate),source_NCLNLSR_loan_weighted_pct=str(noncurrent),
        annualized_NTLNLS_to_endpoint_loans_proxy_pct=str(nco_proxy),
        source_definitions=dict(NCLNLSR='financial_definitions.yaml:12355, noncurrent loans / gross loans',
                                NTLNLS='financial_definitions.yaml:3714, total loan and lease net charge-offs'),
        source_code_evidence='trajectory.jsonl retains the Write tool call for build_deliverables.py; its line 50 uses np.average(g.NCLNLSR, weights=g.LNLSGR) for NCO Rate. Solver deleted the temporary script before final delivery.',
        limitation='Endpoint-loan denominator is a disclosed proxy, not an exact average-loan rate. The wrong source-field identity is independently established regardless of the acceptable NCO denominator alternative.',
        assigned_score=None)
    dest=ROOT/'diagnostics/fin003-gpt-01';dest.mkdir(parents=True,exist_ok=True)
    (dest/'metric_mapping.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':main()
